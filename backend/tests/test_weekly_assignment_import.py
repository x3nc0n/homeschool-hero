from __future__ import annotations

import asyncio
import csv
from datetime import date, datetime
from io import BytesIO, StringIO

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from backend.database import AsyncSessionLocal
from backend.models import Assignment, LessonPlan, Student, Subject
from backend.models.curriculum import CurriculumLesson, CurriculumPackage, CurriculumUnit, LessonResource, Resource
from backend.schemas.bulk_assignment_import import ParsedAssignmentCandidate, ParsedAssignmentDocument, WeeklyImportItem
from backend.services.bulk_assignment_import import BulkAssignmentImportService
from backend.services.weekly_assignment_import import WEEKLY_HEADERS
from tests.contracts import (
    AUTH, CALENDAR, CURRICULUM, STUDENTS, SUBJECTS,
    curriculum_package_payload, school_year_payload, student_payload, subject_payload,
)
from tests.helpers import response_id, sync_csrf_header

IMPORTS = '/api/assignment-import-sessions'
CONTEXT = {'planning_mode': 'weekly', 'week_start': '2026-10-19'}


async def seed(client):
    student = await client.post(STUDENTS['collection'], json=student_payload('Lucas'))
    assert student.status_code == 201, student.text
    year = await client.post(CALENDAR['school_years'], json=school_year_payload(
        name='2026-2027', start_date='2026-08-01', end_date='2027-06-30',
    ))
    assert year.status_code == 201, year.text
    return response_id(student.json()), response_id(year.json())


def row(**changes):
    result = dict(student='Lucas', grade='7', subject='Spelling', unit='Unit 1', lesson='Lesson 6',
                  instructions='Page 25', day='Tuesday', date='')
    result.update(changes)
    return result


def document(rows, format='csv'):
    if format == 'csv':
        stream = StringIO(newline='')
        writer = csv.DictWriter(stream, fieldnames=WEEKLY_HEADERS)
        writer.writeheader()
        writer.writerows(rows)
        return ('weekly.csv', stream.getvalue().encode(), 'text/csv')
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Plans'
    sheet.append(WEEKLY_HEADERS)
    for item in rows:
        sheet.append([item.get(header) for header in WEEKLY_HEADERS])
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return ('weekly.xlsx', stream.getvalue(),
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


async def draft(client, year, rows, format='csv'):
    response = await client.post(IMPORTS, data={**CONTEXT, 'school_year_id': str(year)},
                                 files={'file': document(rows, format)})
    assert response.status_code == 201, response.text
    response = await client.get(f'{IMPORTS}/{response.json()["id"]}')
    assert response.status_code == 200, response.text
    return response.json()


async def confirm(client, payload):
    return await client.post(f'{IMPORTS}/{payload["id"]}/confirm', json={'client_revision': payload['revision']})


def acceptance_rows():
    rows = [row(instructions=f'Page {page}', day=day)
            for page, day in zip(range(25, 29), ('Tuesday', 'Wednesday', 'Thursday', 'Friday'))]
    for lesson, pages, worksheets, day in (
        ('Lesson 4.4', '99-102', '71-72', 'Tuesday'),
        ('Lesson 4.5', '102-105', '73-74', 'Wednesday'),
        ('Lesson 4.6', '105-107', '75-77', 'Thursday'),
    ):
        rows.append(row(subject='Math', unit='Unit 4', lesson=lesson,
                        instructions=f'Textbook {pages}; worksheet {worksheets}', day=day))
    rows.append(row(subject='Math', unit='Unit 4', lesson='Quiz 3', instructions='Quiz 3', day='Friday', category='quiz'))
    rows.append(row(subject='History', unit='Chapter 8', lesson='Chapter 8', instructions='Pages 70-85', day='Friday'))
    for unit, lesson, pages, day in (
        ('Unit 4', '4.8-4.9', '55-57', 'Tuesday'),
        ('Unit 4', '4.10-4.12', '58-61', 'Wednesday'),
        ('Unit 4', 'Chapter 4 Review', '62-63', 'Thursday'),
        ('Unit 5', '5.1-5.3', '65-68', 'Friday'),
    ):
        rows.append(row(subject='Science', unit=unit, lesson=lesson, instructions=f'Pages {pages}', day=day))
    return rows


@pytest.mark.asyncio
@pytest.mark.parametrize('format', ['csv', 'xlsx'])
async def test_lucas_acceptance_creates_linked_records_and_repeat_skips(authorized_client, monkeypatch, format):
    monkeypatch.setattr('backend.config.settings.ai_import_enabled', False)
    student_id, year = await seed(authorized_client)
    payload = await draft(authorized_client, year, acceptance_rows(), format)
    assert payload['parse_method'] == 'weekly_template'
    assert payload['weekly_context'] == {'week_start': '2026-10-19', 'school_year_id': year}
    assert payload['summary'] == {'total': 13, 'ready': 13, 'needs_clarification': 0, 'invalid': 0}
    assert len({item['weekly']['activity_name'] for item in payload['items'][:4]}) == 4
    assert payload['items'][0]['weekly']['resolution']['subject'] == 'new'
    async with AsyncSessionLocal() as db:
        assert await db.scalar(select(func.count()).select_from(Assignment)) == 0
        assert await db.scalar(select(func.count()).select_from(CurriculumPackage)) == 0
    response = await confirm(authorized_client, payload)
    assert response.status_code == 201, response.text
    result = response.json()
    assert result['assignment_count'] == result['lesson_plan_count'] == 13
    assert result['curriculum_counts']['subjects']['created'] == 4
    assert result['curriculum_counts']['activities']['created'] == 13
    async with AsyncSessionLocal() as db:
        subjects = (await db.execute(select(Subject))).scalars().all()
        assert {item.name for item in subjects} == {
            'Spelling - 7th Grade', 'Math - 7th Grade', 'History - 7th Grade', 'Science - 7th Grade',
        }
        assignments = (await db.execute(select(Assignment))).scalars().all()
        assert len(assignments) == 13
        assert len({assignment.lesson_plan_id for assignment in assignments}) == 13
        plans = (await db.execute(select(LessonPlan))).scalars().all()
        assert {plan.student_id for plan in plans} == {student_id}
        assert {plan.school_year_id for plan in plans} == {year}
        assert {plan.target_date.weekday() for plan in plans} == {1, 2, 3, 4}
    repeat = await draft(authorized_client, year, acceptance_rows(), format)
    assert all(item['weekly']['resolution']['duplicate'] for item in repeat['items'])
    response = await confirm(authorized_client, repeat)
    assert response.status_code == 201, response.text
    assert response.json()['assignment_count'] == response.json()['lesson_plan_count'] == 0
    assert len(response.json()['skipped_item_ids']) == 13
    assert response.json()['skipped_count'] == len(response.json()['reused_lesson_plan_ids']) == 13
    assert len(response.json()['reused_assignment_ids']) == 13


@pytest.mark.asyncio
@pytest.mark.parametrize('format', ['csv', 'xlsx'])
async def test_template_download_roundtrip_and_guidance(authorized_client, monkeypatch, backend_module, format):
    monkeypatch.setattr('backend.config.settings.ai_import_enabled', False)
    _, year = await seed(authorized_client)
    response = await authorized_client.get(f'{IMPORTS}/template', params={'format': format})
    assert response.status_code == 200
    assert f'weekly-plan-template.{format}' in response.headers['content-disposition']
    mime = 'text/csv' if format == 'csv' else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    response = await authorized_client.post(IMPORTS, data={**CONTEXT, 'school_year_id': str(year)},
                                            files={'file': (f'weekly.{format}', response.content, mime)})
    assert response.status_code == 201, response.text
    assert response.json()['summary']['ready'] == 1
    schema = backend_module.app.openapi()
    assert 'YYYY-MM-DD' in schema['paths'][f'{IMPORTS}/template']['get']['description']


@pytest.mark.asyncio
@pytest.mark.parametrize('changes,error', [
    ({'date': '2026-10-21', 'day': 'Tuesday'}, 'date and day conflict'),
    ({'date': '2026-10-26', 'day': ''}, 'outside the selected week'),
    ({'day': 'someday'}, 'weekday name'),
    ({'day': '', 'date': ''}, 'weekly.scheduled_date'),
    ({'student': 'Not a real student'}, 'targets'),
    ({'grade': ''}, 'weekly.grade'),
])
async def test_dates_and_missing_fields_require_correction(authorized_client, changes, error):
    _, year = await seed(authorized_client)
    payload = await draft(authorized_client, year, [row(**changes)])
    assert payload['status'] == 'needs_clarification'
    assert error in str(payload['items'][0])
    response = await confirm(authorized_client, payload)
    assert response.status_code == 409
    async with AsyncSessionLocal() as db:
        assert await db.scalar(select(func.count()).select_from(Student)) == 1
        assert await db.scalar(select(func.count()).select_from(Assignment)) == 0


@pytest.mark.asyncio
@pytest.mark.parametrize('header,body', [
    ('student,grade,subject,unit,lesson,instructions,date,student', 'Lucas,7,Math,U,L,Read,2026-10-20,Lucas'),
    ('student,grade,subject,unit,lesson,date', 'Lucas,7,Math,U,L,2026-10-20'),
    ('student,grade,subject,unit,lesson,instructions,when', 'Lucas,7,Math,U,L,Read,Tuesday'),
    ('student,grade,subject,unit,lesson,instructions,date', 'Lucas,7,Math,U,L,Read,10/20/26'),
    ('student,grade,subject,unit,lesson,instructions,day,category', 'Lucas,7,Math,U,L,Read,Tuesday,wrong'),
    ('student,grade,subject,unit,lesson,instructions,day,estimated_minutes', 'Lucas,7,Math,U,L,Read,Tuesday,-1'),
    ('student,grade,subject,unit,lesson,instructions,day,resource_url', 'Lucas,7,Math,U,L,Read,Tuesday,javascript:alert(1)'),
])
async def test_malformed_canonical_templates_never_fall_back_to_ai(authorized_client, monkeypatch, header, body):
    _, year = await seed(authorized_client)
    monkeypatch.setattr('backend.config.settings.ai_import_enabled', False)
    response = await authorized_client.post(
        IMPORTS, data={**CONTEXT, 'school_year_id': str(year)},
        files={'file': ('weekly.csv', f'{header}\n{body}\n'.encode(), 'text/csv')},
    )
    assert response.status_code == 400, response.text


@pytest.mark.asyncio
async def test_xlsx_native_date_and_formula_protection(authorized_client):
    _, year = await seed(authorized_client)
    payload = await draft(authorized_client, year, [row(date=datetime(2026, 10, 20))], 'xlsx')
    assert payload['summary']['ready'] == 1
    filename, content, mime = document([row()], 'xlsx')
    workbook = load_workbook(BytesIO(content))
    workbook.active['A2'] = '=CONCAT("Lu","cas")'
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    response = await authorized_client.post(IMPORTS, data={**CONTEXT, 'school_year_id': str(year)},
                                            files={'file': (filename, stream.getvalue(), mime)})
    assert response.status_code == 400
    assert 'formula' in response.text


@pytest.mark.asyncio
async def test_duplicate_rows_skip_and_existing_date_conflict_is_not_moved(authorized_client):
    _, year = await seed(authorized_client)
    payload = await draft(authorized_client, year, [row(), row()])
    assert payload['items'][1]['weekly']['resolution']['duplicate_of'] == payload['items'][0]['client_item_id']
    response = await confirm(authorized_client, payload)
    assert response.status_code == 201, response.text
    assert response.json()['assignment_count'] == response.json()['lesson_plan_count'] == 1
    conflicting = await draft(authorized_client, year, [row(day='Wednesday')])
    assert conflicting['items'][0]['weekly']['resolution']['lesson_plan'] == 'conflict'
    assert (await confirm(authorized_client, conflicting)).status_code == 409
    async with AsyncSessionLocal() as db:
        assert (await db.scalar(select(LessonPlan))).target_date == date(2026, 10, 20)
        assert await db.scalar(select(func.count()).select_from(Assignment)) == 1


@pytest.mark.asyncio
async def test_new_draft_conflicting_dates_block_before_any_creation(authorized_client):
    _, year = await seed(authorized_client)
    payload = await draft(authorized_client, year, [row(), row(day='Wednesday')])
    assert payload['summary']['invalid'] == 1
    assert (await confirm(authorized_client, payload)).status_code == 409


@pytest.mark.asyncio
async def test_ambiguous_curriculum_can_be_selected_and_resources_are_linked(authorized_client):
    _, year = await seed(authorized_client)
    subject = await authorized_client.post(SUBJECTS['collection'], json=subject_payload(name='Spelling - 7th Grade'))
    packages = []
    for name in ('Spelling A', 'Spelling B'):
        response = await authorized_client.post(CURRICULUM['packages'],
                                               json=curriculum_package_payload(year, response_id(subject.json()), name=name))
        assert response.status_code == 201, response.text
        packages.append(response_id(response.json()))
    payload = await draft(authorized_client, year, [row(resource_url='https://example.org/page25')])
    assert payload['items'][0]['weekly']['resolution']['package'] == 'ambiguous'
    question = next(question for question in payload['questions'] if question['field'] == 'weekly.package_id')
    response = await authorized_client.patch(f'{IMPORTS}/{payload["id"]}', json={
        'answers': [{'question_id': question['id'], 'value': packages[1]}],
    })
    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload['status'] == 'ready'
    assert payload['items'][0]['weekly']['resolution']['package_id'] == packages[1]
    assert (await confirm(authorized_client, payload)).status_code == 201
    async with AsyncSessionLocal() as db:
        assert await db.scalar(select(func.count()).select_from(Resource)) == 1
        assert await db.scalar(select(func.count()).select_from(LessonResource)) == 1
        assert (await db.scalar(select(CurriculumUnit))).package_id == packages[1]


@pytest.mark.asyncio
async def test_paste_context_ai_weekly_shape_and_editing(authorized_client, monkeypatch):
    student_id, year = await seed(authorized_client)
    monkeypatch.setattr(BulkAssignmentImportService, 'ensure_ai_configured', lambda self: None)
    seen = {}

    def capture(session_id, source, defaults):
        seen.update(session_id=session_id, source=source, defaults=defaults)

    monkeypatch.setattr('backend.routers.bulk_assignment_import.schedule_assignment_import_processing', capture)
    response = await authorized_client.post(IMPORTS, data={
        **CONTEXT, 'school_year_id': str(year), 'text': 'Lucas 7th grade Spelling Lesson 6 page 25 Tuesday',
    })
    assert response.status_code == 201, response.text
    assert response.json()['status'] == 'processing'
    service = BulkAssignmentImportService()
    request = service._build_request_payload(seen['source'].extracted, model='test-model')
    assert 'WEEKLY MODE' in request['messages'][0]['content']
    assert '2026-10-19' in request['messages'][1]['content']

    async def parsed(_):
        return ParsedAssignmentDocument(assignments=[ParsedAssignmentCandidate(
            title='Lesson 6', description='Page 25', weekly=WeeklyImportItem(**row()),
        )])

    monkeypatch.setattr(service, 'parse_with_ai', parsed)
    async with AsyncSessionLocal() as db:
        await service.complete_processing_session(db, session_id=seen['session_id'],
                                                  source=seen['source'], defaults=seen['defaults'])
    payload = (await authorized_client.get(f'{IMPORTS}/{seen["session_id"]}')).json()
    item = payload['items'][0]
    item['title'] = 'Edited assignment'
    item['description'] = 'Edited instructions for assignment'
    item['weekly']['day'] = 'Wednesday'
    item['weekly']['scheduled_date'] = '2026-10-21'
    item['targets'] = [{'student_id': student_id}]
    payload = (await authorized_client.patch(f'{IMPORTS}/{payload["id"]}', json={'items': [item]})).json()
    assert payload['items'][0]['due_date'].startswith('2026-10-21')
    response = await confirm(authorized_client, payload)
    assert response.status_code == 201, response.text
    async with AsyncSessionLocal() as db:
        assignment = await db.scalar(select(Assignment))
        assert assignment.title == 'Edited assignment'
        assert assignment.description == 'Edited instructions for assignment'
        assert (await db.scalar(select(LessonPlan))).target_date == date(2026, 10, 21)


@pytest.mark.asyncio
async def test_revision_double_confirmation_and_transaction_rollback(authorized_client, monkeypatch):
    _, year = await seed(authorized_client)
    payload = await draft(authorized_client, year, [row()])
    response = await authorized_client.post(f'{IMPORTS}/{payload["id"]}/confirm', json={'client_revision': 999})
    assert response.status_code == 409
    original = BulkAssignmentImportService.confirm

    async def fail_after_creation(self, *args, **kwargs):
        await original(self, *args, **kwargs)
        from fastapi import HTTPException
        raise HTTPException(409, 'Injected failure after linked creation')

    monkeypatch.setattr(BulkAssignmentImportService, 'confirm', fail_after_creation)
    assert (await confirm(authorized_client, payload)).status_code == 409
    async with AsyncSessionLocal() as db:
        for model in (Assignment, LessonPlan, CurriculumLesson, CurriculumUnit, CurriculumPackage, Subject):
            assert await db.scalar(select(func.count()).select_from(model)) == 0
    monkeypatch.setattr(BulkAssignmentImportService, 'confirm', original)
    responses = await asyncio.gather(confirm(authorized_client, payload), confirm(authorized_client, payload))
    assert sorted(response.status_code for response in responses) == [201, 409]
    async with AsyncSessionLocal() as db:
        assert await db.scalar(select(func.count()).select_from(Assignment)) == 1


@pytest.mark.asyncio
async def test_explicit_week_year_and_single_source_are_required(authorized_client):
    _, year = await seed(authorized_client)
    for fields in (
        {'planning_mode': 'weekly'},
        {**CONTEXT, 'school_year_id': '999999'},
        {**CONTEXT, 'school_year_id': str(year), 'week_start': '2026-10-20'},
        {**CONTEXT, 'school_year_id': str(year), 'text': 'also a text input'},
    ):
        response = await authorized_client.post(IMPORTS, data=fields, files={'file': document([row()])})
        assert response.status_code in {400, 404}, response.text


def grid_document():
    workbook = Workbook()
    sheet = workbook.active
    sheet['A1'] = 'Lucas Week 9 October 26-30'
    sheet['A4'] = 'Subject'
    for column, day in zip((4, 6, 8, 10, 12), ('Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday')):
        sheet.cell(4, column, day)
        sheet.cell(4, column + 1, '✓')
    subject_rows = {'Spelling': 6, 'Math': 12, 'History': 14, 'Science': 18}
    for subject, index in subject_rows.items():
        sheet.cell(index, 1, subject)
    for subject, index in (('Literature', 8), ('Grammar', 10), ('Religion', 16)):
        sheet.cell(index, 1, subject)
    weekdays = {'Tuesday': 6, 'Wednesday': 8, 'Thursday': 10, 'Friday': 12}
    for item in acceptance_rows():
        subject = item['subject']
        text = item['lesson'] if subject in {'Spelling', 'History'} else f'{item["unit"]} — {item["lesson"]}'
        text += f'\n{item["instructions"]}'
        sheet.cell(subject_rows[subject], weekdays[item['day']], text)
    # A checked completion box is not another assignment.
    sheet['G6'] = '✓'
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return ('Lucas Week 9.xlsx', stream.getvalue(),
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@pytest.mark.asyncio
async def test_subject_weekday_grid_preserves_13_assignments_and_clarifies_grade(authorized_client, monkeypatch):
    _, year = await seed(authorized_client)
    monkeypatch.setattr('backend.config.settings.ai_import_enabled', False)
    response = await authorized_client.post(
        IMPORTS, data={**CONTEXT, 'school_year_id': str(year), 'week_start': '2026-10-26'},
        files={'file': grid_document()},
    )
    assert response.status_code == 201, response.text
    payload = (await authorized_client.get(f'{IMPORTS}/{response.json()["id"]}')).json()
    assert payload['parse_method'] == 'weekly_grid'
    assert payload['summary']['total'] == 13
    assert payload['summary']['needs_clarification'] == 13
    assert {item['weekly']['subject'] for item in payload['items']} == {'Spelling', 'Math', 'History', 'Science'}
    question = next(question for question in payload['questions'] if question['field'] == 'weekly.grade')
    response = await authorized_client.patch(f'{IMPORTS}/{payload["id"]}', json={
        'answers': [{'question_id': question['id'], 'value': '7'}],
    })
    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload['summary']['ready'] == 13
    assert {item['weekly']['scheduled_date'] for item in payload['items']} == {
        '2026-10-27', '2026-10-28', '2026-10-29', '2026-10-30',
    }
    assert all(item['source_excerpt'] for item in payload['items'])
    response = await confirm(authorized_client, payload)
    assert response.status_code == 201, response.text
    assert response.json()['assignment_count'] == response.json()['lesson_plan_count'] == 13


@pytest.mark.asyncio
async def test_grid_title_conflict_blocks_and_explicit_grade_default_is_supported(authorized_client, monkeypatch):
    _, year = await seed(authorized_client)
    monkeypatch.setattr('backend.config.settings.ai_import_enabled', False)
    for week, expected_ready in (('2026-10-19', 0), ('2026-10-26', 13)):
        response = await authorized_client.post(
            IMPORTS, data={**CONTEXT, 'school_year_id': str(year), 'week_start': week, 'defaults': '{"grade":"7"}'},
            files={'file': grid_document()},
        )
        assert response.status_code == 201, response.text
        payload = (await authorized_client.get(f'{IMPORTS}/{response.json()["id"]}')).json()
        assert payload['summary']['ready'] == expected_ready
        if not expected_ready:
            assert all('conflicts with selected week' in str(item['errors']) for item in payload['items'])
            assert (await confirm(authorized_client, payload)).status_code == 409


@pytest.mark.asyncio
async def test_weekly_matches_and_explicit_package_choices_stay_in_family(
    authorized_client, secondary_client, create_family_user,
):
    _, year = await seed(authorized_client)
    payload = await draft(authorized_client, year, [row()])
    other = await create_family_user(
        family_name='Other weekly family', email='other-weekly@example.com', password='strongpass456',
        role='parent', is_owner=True,
    )
    login = await secondary_client.post(AUTH['login'], json={
        'email': other['email'], 'password': other['password'], 'family_id': other['family_id'],
    })
    assert login.status_code == 200, login.text
    sync_csrf_header(secondary_client)
    _, foreign_year = await seed(secondary_client)
    subject = await secondary_client.post(SUBJECTS['collection'], json=subject_payload(name='Spelling - 7th Grade'))
    package = await secondary_client.post(CURRICULUM['packages'], json=curriculum_package_payload(
        foreign_year, response_id(subject.json()), name='Foreign curriculum',
    ))
    assert package.status_code == 201, package.text
    for method, suffix, body in (('get', '', None), ('patch', '', {'items': []}), ('post', '/confirm', {})):
        kwargs = {'json': body} if body is not None else {}
        response = await getattr(secondary_client, method)(f'{IMPORTS}/{payload["id"]}{suffix}', **kwargs)
        assert response.status_code == 404
    response = await authorized_client.post(
        IMPORTS, data={**CONTEXT, 'school_year_id': str(foreign_year)}, files={'file': document([row()])},
    )
    assert response.status_code == 404
    item = payload['items'][0]
    item['weekly']['package_id'] = response_id(package.json())
    response = await authorized_client.patch(f'{IMPORTS}/{payload["id"]}', json={'items': [item]})
    assert response.status_code == 200, response.text
    assert 'Selected curriculum does not match' in str(response.json()['items'][0]['errors'])
    assert (await confirm(authorized_client, response.json())).status_code == 409


@pytest.mark.asyncio
async def test_two_distinct_sessions_confirm_same_work_once(authorized_client):
    _, year = await seed(authorized_client)
    first = await draft(authorized_client, year, [row()])
    second = await draft(authorized_client, year, [row()])
    responses = await asyncio.gather(confirm(authorized_client, first), confirm(authorized_client, second))
    assert all(response.status_code == 201 for response in responses), [response.text for response in responses]
    assert sorted(response.json()['assignment_count'] for response in responses) == [0, 1]
    async with AsyncSessionLocal() as db:
        for model in (Assignment, LessonPlan, CurriculumLesson, CurriculumUnit, CurriculumPackage, Subject):
            assert await db.scalar(select(func.count()).select_from(model)) == 1


@pytest.mark.asyncio
async def test_patch_revision_and_plain_student_subject_edits_re_resolve(authorized_client):
    _, year = await seed(authorized_client)
    subject = await authorized_client.post(SUBJECTS['collection'], json=subject_payload(name='Spelling - 7th Grade'))
    assert subject.status_code == 201
    payload = await draft(authorized_client, year, [row()])
    assert payload['items'][0]['subject_id'] == response_id(subject.json())
    response = await authorized_client.patch(f'{IMPORTS}/{payload["id"]}', json={'client_revision': 999, 'items': []})
    assert response.status_code == 409
    second_student = await authorized_client.post(STUDENTS['collection'], json=student_payload('Grace'))
    item = payload['items'][0]
    item['weekly']['subject'] = 'Math'
    item['weekly']['student'] = 'Grace'
    response = await authorized_client.patch(f'{IMPORTS}/{payload["id"]}', json={'items': [item]})
    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload['status'] == 'ready', payload
    assert payload['items'][0]['targets'][0]['student_id'] == response_id(second_student.json())
    assert payload['items'][0]['weekly']['resolution']['subject_name'] == 'Math - 7th Grade'
    response = await confirm(authorized_client, payload)
    assert response.status_code == 201, response.text


@pytest.mark.asyncio
async def test_confirm_selected_ready_rows_skips_invalid_and_missing_rows(authorized_client):
    _, year = await seed(authorized_client)
    payload = await draft(authorized_client, year, [
        row(),
        row(instructions='Page 26', day='Wednesday', date='2026-10-20'),
        row(instructions='Page 27', student='Unknown student'),
    ])
    assert payload['status'] == 'needs_clarification'
    assert payload['summary']['ready'] == 1
    response = await confirm(authorized_client, payload)
    assert response.status_code == 409
    response = await authorized_client.post(f'{IMPORTS}/{payload["id"]}/confirm', json={
        'client_revision': payload['revision'],
        'item_ids': [payload['items'][0]['client_item_id']],
    })
    assert response.status_code == 201, response.text
    assert response.json()['assignment_count'] == response.json()['lesson_plan_count'] == 1
    assert set(response.json()['skipped_item_ids']) == {
        item['client_item_id'] for item in payload['items'][1:]
    }


@pytest.mark.asyncio
async def test_weekly_subset_still_revalidates_selected_conflicts_and_rejects_empty_selection(authorized_client):
    _, year = await seed(authorized_client)
    payload = await draft(authorized_client, year, [row(), row(instructions='Page 26', day='invalid')])
    for ids in ([], [payload['items'][1]['client_item_id']], ['unknown_item']):
        response = await authorized_client.post(f'{IMPORTS}/{payload["id"]}/confirm', json={
            'client_revision': payload['revision'], 'item_ids': ids,
        })
        assert response.status_code in {400, 409}, response.text
    async with AsyncSessionLocal() as db:
        assert await db.scalar(select(func.count()).select_from(Assignment)) == 0
