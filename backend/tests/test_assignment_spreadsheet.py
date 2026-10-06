from __future__ import annotations

from datetime import datetime
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile

import pytest
from fastapi import UploadFile
from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904
from starlette.datastructures import Headers

from backend.services import assignment_spreadsheet as spreadsheet
from backend.services.bulk_assignment_import import BulkAssignmentImportError, BulkAssignmentImportService
from backend.services.curriculum_ai_import import AIImportError


def workbook_bytes(workbook: Workbook | None = None) -> bytes:
    buffer = BytesIO()
    (workbook or Workbook()).save(buffer)
    return buffer.getvalue()


def replace_zip_entry(payload: bytes, name: str, transform) -> bytes:
    result = BytesIO()
    with ZipFile(BytesIO(payload)) as source, ZipFile(result, 'w', ZIP_DEFLATED) as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            target.writestr(entry.filename, transform(data) if entry.filename == name else data)
    return result.getvalue()


def grid_workbook() -> Workbook:
    workbook = Workbook()
    week = workbook.active
    week.title = 'Week 1'
    week.merge_cells('A1:C1')
    week['A1'] = 'Ada — Math'
    week.append(['Subject', 'Monday', 'Tuesday'])
    week.append(['Date', datetime(2026, 10, 12), datetime(2026, 10, 13)])
    week.append(['Math', 'Fractions\npage 2', 'Read chapter 3'])
    workbook.create_sheet('Reading')['B2'] = 'Read "Little Women"'
    workbook.create_sheet('Hidden notes')['A1'] = 'Review notes'
    workbook['Hidden notes'].sheet_state = 'hidden'
    return workbook


def extract(payload: bytes, max_input_chars: int = 50_000):
    return spreadsheet.extract_assignment_spreadsheet(payload, max_input_chars=max_input_chars)


def test_multisheet_grid_dates_merged_headers_and_hidden_sheets():
    text, warnings = extract(workbook_bytes(grid_workbook()))
    assert warnings == []
    assert 'Sheet: "Week 1"' in text
    assert 'Merged range: A1:C1; header/value is at A1' in text
    assert 'Row 2: A2="Subject" | B2="Monday" | C2="Tuesday"' in text
    assert 'B3="2026-10-12T00:00:00" | C3="2026-10-13T00:00:00"' in text
    assert 'Row 4: A4="Math" | B4="Fractions\\npage 2" | C4="Read chapter 3"' in text
    assert 'Sheet: "Reading"' in text and 'Row 2: B2="Read \\"Little Women\\""' in text
    assert 'Sheet: "Hidden notes" (state: hidden)' in text
    assert 'A1="Review notes"' in text


def test_dates_respect_workbook_1904_epoch():
    workbook = Workbook()
    workbook.epoch = CALENDAR_MAC_1904
    workbook.active['A1'] = datetime(2026, 10, 13, 14, 30)
    text, _ = extract(workbook_bytes(workbook))
    assert 'A1="2026-10-13T14:30:00"' in text


def test_formulas_missing_cached_values_are_explicit_and_not_evaluated():
    workbook = Workbook()
    workbook.active['A1'] = '=1+2'
    workbook.active['B1'] = '=HYPERLINK("https://example.invalid","external")'
    text, warnings = extract(workbook_bytes(workbook))
    assert 'A1=[formula; cached value unavailable; do not infer a value]' in text
    assert 'B1=[formula; cached value unavailable; do not infer a value]' in text
    assert '=1+2' not in text and 'https://' not in text
    assert '2 formula cell(s) were not evaluated' in warnings[0]
    assert '2 formula cell(s) have no cached value' in warnings[1]


@pytest.mark.parametrize(
    ('cell_type', 'cached_value', 'expected'),
    [('', '3', '3'), (' t="str"', 'cached text', '"cached text"'), (' t="b"', '0', 'false')],
)
def test_saved_formula_results_used_with_staleness_warning(cell_type, cached_value, expected):
    workbook = Workbook()
    workbook.active['A1'] = '=1+2'
    payload = replace_zip_entry(
        workbook_bytes(workbook), 'xl/worksheets/sheet1.xml',
        lambda data: data.replace(
            b'<c r="A1"><f>1+2</f><v></v></c>',
            f'<c r="A1"{cell_type}><f>1+2</f><v>{cached_value}</v></c>'.encode(),
        ),
    )
    text, warnings = extract(payload)
    assert f'A1={expected} [cached formula value; may be stale]' in text
    assert len(warnings) == 1 and 'may be stale' in warnings[0]


@pytest.mark.parametrize('payload', [b'not a zip', b'PK\x03\x04broken', spreadsheet._OLE_HEADER + b'encrypted'])
def test_malformed_or_encrypted_workbooks_fail_explicitly(payload):
    with pytest.raises(AIImportError, match='malformed|Password-protected'):
        extract(payload)


def test_empty_and_whitespace_only_workbooks_rejected():
    workbook = Workbook()
    for value in (None, ' \n\t'):
        workbook.active['A1'] = value
        with pytest.raises(AIImportError, match='empty or contains no readable'):
            extract(workbook_bytes(workbook))


@pytest.mark.parametrize(
    ('limit', 'value', 'message'),
    [
        ('MAX_ZIP_ENTRIES', 1, 'ZIP entries'),
        ('MAX_ZIP_ENTRY_BYTES', 1, 'expanded bytes per ZIP entry'),
        ('MAX_ZIP_TOTAL_BYTES', 1, 'total expanded ZIP bytes'),
        ('MAX_COMPRESSION_RATIO', 1, 'compression ratio'),
        ('MAX_XML_ELEMENTS', 1, 'XML elements'),
        ('MAX_SHEETS', 1, 'sheets'),
        ('MAX_TOTAL_ROWS', 1, 'total rows'),
        ('MAX_CELLS', 1, 'cells'),
        ('MAX_MERGES', 0, 'merged ranges'),
        ('MAX_TEXT_CHARS', 1, 'extracted text characters'),
    ],
)
def test_resource_limits_fail_before_ai(monkeypatch, limit, value, message):
    monkeypatch.setattr(spreadsheet, limit, value)
    with pytest.raises(AIImportError, match=message):
        extract(workbook_bytes(grid_workbook()))


@pytest.mark.parametrize('coordinate', ['A5001', 'CW1', 'XFD1048576'])
def test_sparse_huge_coordinates_rejected(coordinate):
    workbook = Workbook()
    workbook.active[coordinate] = 'Huge grid'
    with pytest.raises(AIImportError, match='rows / 100 columns'):
        extract(workbook_bytes(workbook))


def test_merged_range_extents_bounded():
    workbook = Workbook()
    workbook.active['A1'] = 'Header'
    payload = replace_zip_entry(
        workbook_bytes(workbook), 'xl/worksheets/sheet1.xml',
        lambda data: data.replace(b'</worksheet>', b'<mergeCells><mergeCell ref="A1:XFD1048576"/></mergeCells></worksheet>'),
    )
    with pytest.raises(AIImportError, match='rows / 100 columns'):
        extract(payload)


def test_sparse_grid_scanned_cells_bounded(monkeypatch):
    workbook = Workbook()
    for row in range(1, 5):
        workbook.active[f'C{row}'] = 'Sparse grid'
    monkeypatch.setattr(spreadsheet, 'MAX_CELLS', 10)
    with pytest.raises(AIImportError, match='scanned cells'):
        extract(workbook_bytes(workbook))


def test_inflated_dimension_does_not_expand_grid():
    workbook = Workbook()
    workbook.active['A1'] = 'Assignment'
    payload = replace_zip_entry(
        workbook_bytes(workbook), 'xl/worksheets/sheet1.xml',
        lambda data: data.replace(b'ref="A1:A1"', b'ref="A1:XFD1048576"'),
    )
    text, _ = extract(payload)
    assert 'Row 1: A1="Assignment"' in text


def test_xml_entities_rejected():
    payload = replace_zip_entry(
        workbook_bytes(grid_workbook()), 'xl/worksheets/sheet1.xml',
        lambda data: b'<!DOCTYPE worksheet [<!ENTITY entity "unsafe">]>' + data,
    )
    with pytest.raises(AIImportError, match='malformed'):
        extract(payload)


def test_inconsistent_row_and_cell_coordinates_rejected():
    workbook = Workbook()
    workbook.active['A1'] = 'Assignment'
    payload = replace_zip_entry(
        workbook_bytes(workbook), 'xl/worksheets/sheet1.xml',
        lambda data: data.replace(b'<c r="A1"', b'<c r="A2"'),
    )
    with pytest.raises(AIImportError, match='coordinates are inconsistent'):
        extract(payload)


def test_macro_enabled_workbook_renamed_xlsx_rejected():
    payload = replace_zip_entry(
        workbook_bytes(grid_workbook()), '[Content_Types].xml',
        lambda data: data.replace(
            spreadsheet._WORKBOOK_TYPE.encode(),
            b'application/vnd.ms-excel.sheet.macroEnabled.main+xml',
        ),
    )
    with pytest.raises(AIImportError, match='Unsupported Excel workbook format'):
        extract(payload)


def test_duplicate_zip_members_rejected():
    result = BytesIO()
    with ZipFile(result, 'w', ZIP_STORED) as archive:
        archive.writestr('duplicate.xml', '<a/>')
        with pytest.warns(UserWarning):
            archive.writestr('duplicate.xml', '<a/>')
    with pytest.raises(AIImportError, match='duplicate ZIP entries'):
        extract(result.getvalue())


def test_encrypted_zip_members_rejected_before_reading():
    payload = bytearray(workbook_bytes(grid_workbook()))
    local_header = payload.index(b'PK\x03\x04')
    central_header = payload.index(b'PK\x01\x02')
    payload[local_header + 6] |= 1
    payload[central_header + 8] |= 1
    with pytest.raises(AIImportError, match='Password-protected'):
        extract(bytes(payload))


def test_zip_bomb_compression_rejected_with_default_limits():
    result = BytesIO()
    with ZipFile(result, 'w', ZIP_DEFLATED) as archive:
        archive.writestr('large.xml', b' ' * (1024 * 1024))
    with pytest.raises(AIImportError, match='compression ratio'):
        extract(result.getvalue())


def test_non_xml_worksheet_relationship_rejected():
    payload = replace_zip_entry(
        workbook_bytes(grid_workbook()), 'xl/_rels/workbook.xml.rels',
        lambda data: data.replace(b'sheet1.xml', b'sheet1.dat'),
    )
    with pytest.raises(AIImportError, match='worksheet relationship'):
        extract(payload)


def test_ai_text_limit_rejects_instead_of_silently_truncating_grid():
    with pytest.raises(AIImportError, match='100 extracted text characters'):
        extract(workbook_bytes(grid_workbook()), max_input_chars=100)


@pytest.mark.asyncio
@pytest.mark.parametrize('content_type', [spreadsheet.XLSX_CONTENT_TYPE, 'application/octet-stream', 'application/zip', ''])
async def test_upload_accepts_xlsx_and_preserves_lines(content_type):
    upload = UploadFile(
        BytesIO(workbook_bytes(grid_workbook())), filename='plan.XLSX',
        headers=Headers({'content-type': content_type}),
    )
    source = await BulkAssignmentImportService().extract_upload(upload)
    assert source.content_type == spreadsheet.XLSX_CONTENT_TYPE
    assert '\nRow 2:' in source.extracted.text
    assert source.text_hash


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ('filename', 'content_type'),
    [('plan.xlsx', 'text/plain'), ('plan.txt', spreadsheet.XLSX_CONTENT_TYPE),
     ('plan.xls', spreadsheet.XLSX_CONTENT_TYPE), ('plan.xlsm', spreadsheet.XLSX_CONTENT_TYPE)],
)
async def test_upload_rejects_unsupported_and_mismatched_spreadsheets(filename, content_type):
    upload = UploadFile(
        BytesIO(workbook_bytes(grid_workbook())), filename=filename,
        headers=Headers({'content-type': content_type}),
    )
    with pytest.raises(BulkAssignmentImportError, match='MIME|filename|Unsupported spreadsheet'):
        await BulkAssignmentImportService().extract_upload(upload)
