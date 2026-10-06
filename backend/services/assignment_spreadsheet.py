from __future__ import annotations

import json
from datetime import date, datetime, time, timedelta
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile

from defusedxml.ElementTree import iterparse
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries

from backend.services.curriculum_ai_import import AIImportError

XLSX_CONTENT_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
MAX_ZIP_ENTRIES = 256
MAX_ZIP_ENTRY_BYTES = 8 * 1024 * 1024
MAX_ZIP_TOTAL_BYTES = 20 * 1024 * 1024
MAX_COMPRESSION_RATIO = 200
MAX_XML_ELEMENTS = 400_000
MAX_SHEETS = 20
MAX_ROWS_PER_SHEET = 5_000
MAX_TOTAL_ROWS = 20_000
MAX_COLUMNS = 100
MAX_CELLS = 100_000
MAX_MERGES = 1_000
MAX_TEXT_CHARS = 200_000
_NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
_WORKBOOK_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml'
_OLE_HEADER = b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'


def _limit(message: str) -> None:
    raise AIImportError(f'Excel workbook exceeds the {message} limit. Split the workbook into smaller files.')


def _check_grid(row: int, column: int) -> None:
    if row < 1 or column < 1:
        raise AIImportError('Malformed Excel workbook: invalid cell coordinates.')
    if row > MAX_ROWS_PER_SHEET or column > MAX_COLUMNS:
        _limit(f'{MAX_ROWS_PER_SHEET} rows / {MAX_COLUMNS} columns per sheet')


def _preflight(archive: ZipFile) -> dict[str, list[str]]:
    entries = archive.infolist()
    if len(entries) > MAX_ZIP_ENTRIES:
        _limit(f'{MAX_ZIP_ENTRIES} ZIP entries')
    if len({entry.filename for entry in entries}) != len(entries):
        raise AIImportError('Malformed Excel workbook: duplicate ZIP entries.')
    expanded_bytes = elements = sheets = rows = cells = merges = 0
    workbook_type_found = False
    merged_ranges: dict[str, list[str]] = {}
    for entry in entries:
        if entry.flag_bits & 1:
            raise AIImportError('Password-protected Excel workbooks are not supported. Upload an unencrypted .xlsx file.')
        if entry.compress_type not in {ZIP_STORED, ZIP_DEFLATED}:
            raise AIImportError('Unsupported Excel ZIP compression. Save the workbook again as .xlsx.')
        if entry.file_size > MAX_ZIP_ENTRY_BYTES:
            _limit(f'{MAX_ZIP_ENTRY_BYTES} expanded bytes per ZIP entry')
        expanded_bytes += entry.file_size
        if expanded_bytes > MAX_ZIP_TOTAL_BYTES:
            _limit(f'{MAX_ZIP_TOTAL_BYTES} total expanded ZIP bytes')
        if entry.file_size > max(1, entry.compress_size) * MAX_COMPRESSION_RATIO:
            _limit(f'{MAX_COMPRESSION_RATIO}:1 ZIP compression ratio')
        # Read and verify bounded contents before openpyxl allocates shared strings/styles.
        with archive.open(entry) as source:
            data = source.read(MAX_ZIP_ENTRY_BYTES + 1)
        if len(data) > MAX_ZIP_ENTRY_BYTES or len(data) != entry.file_size:
            _limit(f'{MAX_ZIP_ENTRY_BYTES} expanded bytes per ZIP entry')
        if not entry.filename.lower().endswith(('.xml', '.rels')):
            continue
        ranges: list[str] = []
        depth = sheet_rows = current_row = last_row = last_column = 0
        is_worksheet = False
        for event, element in iterparse(BytesIO(data), events=('start', 'end'), forbid_dtd=True):
            if event == 'start':
                depth += 1
                elements += 1
                if depth > 64 or elements > MAX_XML_ELEMENTS:
                    _limit(f'{MAX_XML_ELEMENTS} XML elements / 64 nesting levels')
                if depth == 1:
                    is_worksheet = element.tag == f'{{{_NS}}}worksheet'
                if element.tag == f'{{{_NS}}}sheet':
                    sheets += 1
                    if sheets > MAX_SHEETS:
                        _limit(f'{MAX_SHEETS} sheets')
                if element.tag.endswith('}Override') and element.attrib.get('PartName') == '/xl/workbook.xml':
                    if element.attrib.get('ContentType') != _WORKBOOK_TYPE:
                        raise AIImportError('Unsupported Excel workbook format. Upload .xlsx, not .xls or .xlsm.')
                    workbook_type_found = True
                if element.tag.endswith('}Override') and element.attrib.get('ContentType', '').endswith('+xml'):
                    if not element.attrib.get('PartName', '').lower().endswith('.xml'):
                        raise AIImportError('Unsupported Excel XML part name. Save the workbook again as .xlsx.')
                if element.tag.endswith('}Relationship'):
                    relationship_type = element.attrib.get('Type', '').rsplit('/', 1)[-1]
                    if relationship_type in {'worksheet', 'styles', 'sharedStrings'}:
                        if element.attrib.get('TargetMode') == 'External' or not element.attrib.get('Target', '').lower().endswith('.xml'):
                            raise AIImportError('Unsupported Excel worksheet relationship. Save the workbook again as .xlsx.')
                if is_worksheet and element.tag == f'{{{_NS}}}row':
                    rows += 1
                    sheet_rows += 1
                    current_row = int(element.attrib['r'])
                    _check_grid(current_row, 1)
                    if current_row <= last_row:
                        raise AIImportError('Malformed Excel workbook: worksheet rows are out of order or duplicated.')
                    last_row = current_row
                    last_column = 0
                    if rows > MAX_TOTAL_ROWS or sheet_rows > MAX_ROWS_PER_SHEET:
                        _limit(f'{MAX_TOTAL_ROWS} total rows / {MAX_ROWS_PER_SHEET} rows per sheet')
                if is_worksheet and element.tag == f'{{{_NS}}}c':
                    cells += 1
                    cell_row, cell_column = coordinate_to_tuple(element.attrib['r'])
                    _check_grid(cell_row, cell_column)
                    if cell_row != current_row or cell_column <= last_column:
                        raise AIImportError('Malformed Excel workbook: worksheet cell coordinates are inconsistent.')
                    last_column = cell_column
                    if cells > MAX_CELLS:
                        _limit(f'{MAX_CELLS} cells')
                if is_worksheet and element.tag == f'{{{_NS}}}mergeCell':
                    merges += 1
                    if merges > MAX_MERGES:
                        _limit(f'{MAX_MERGES} merged ranges')
                    reference = element.attrib['ref']
                    min_column, min_row, max_column, max_row = range_boundaries(reference)
                    _check_grid(min_row, min_column)
                    _check_grid(max_row, max_column)
                    ranges.append(reference)
            else:
                if is_worksheet and element.tag == f'{{{_NS}}}row':
                    current_row = 0
                depth -= 1
                element.clear()
        if is_worksheet:
            merged_ranges[entry.filename] = ranges
    if not workbook_type_found:
        raise AIImportError('Malformed or unsupported Excel workbook. Save it as a standard .xlsx file.')
    return merged_ranges


def _cell_text(value) -> str:
    if isinstance(value, (date, datetime, time)):
        value = value.isoformat()
    elif isinstance(value, timedelta):
        value = str(value)
    return json.dumps(value, ensure_ascii=False, default=str)


def extract_assignment_spreadsheet(payload: bytes, *, max_input_chars: int) -> tuple[str, list[str]]:
    if payload.startswith(_OLE_HEADER):
        raise AIImportError('Password-protected or legacy Excel workbooks are not supported. Upload an unencrypted .xlsx file.')
    workbooks = []
    try:
        with ZipFile(BytesIO(payload)) as archive:
            merged_ranges = _preflight(archive)
        # One stream preserves formula markers; the other reads saved results only.
        # Neither evaluates formulas or fetches external links.
        for data_only in (False, True):
            workbooks.append(load_workbook(BytesIO(payload), read_only=True, data_only=data_only, keep_links=False))
        workbook, cached_workbook = workbooks
        if len(workbook.sheetnames) > MAX_SHEETS:
            _limit(f'{MAX_SHEETS} sheets')
        if len(workbook.worksheets) != len(workbook.sheetnames):
            raise AIImportError('Unsupported Excel sheet type. Use worksheets, not chart sheets.')
        parts: list[str] = []
        text_length = values = formulas = missing_results = scanned_cells = scanned_rows = 0
        max_chars = min(MAX_TEXT_CHARS, max_input_chars)

        def append(part: str) -> None:
            nonlocal text_length
            text_length += len(part) + (1 if parts else 0)
            if text_length > max_chars:
                _limit(f'{max_chars} extracted text characters')
            parts.append(part)

        for sheet, cached_sheet in zip(workbook.worksheets, cached_workbook.worksheets, strict=True):
            append(f'Sheet: {json.dumps(sheet.title, ensure_ascii=False)} (state: {sheet.sheet_state})')
            for reference in merged_ranges.get(sheet._worksheet_path, []):
                append(f'Merged range: {reference}; header/value is at {reference.split(":")[0]} and applies to the entire range.')
            # Ignore inflated dimension metadata; actual coordinates were validated above.
            sheet.reset_dimensions()
            cached_sheet.reset_dimensions()
            for row_index, (row, cached_row) in enumerate(zip(sheet.iter_rows(), cached_sheet.iter_rows(), strict=True), 1):
                scanned_rows += 1
                if row_index > MAX_ROWS_PER_SHEET or scanned_rows > MAX_TOTAL_ROWS:
                    _limit(f'{MAX_TOTAL_ROWS} total rows / {MAX_ROWS_PER_SHEET} rows per sheet')
                scanned_cells += len(row)
                if scanned_cells > MAX_CELLS or len(row) > MAX_COLUMNS:
                    _limit(f'{MAX_CELLS} scanned cells / {MAX_COLUMNS} columns per sheet')
                row_parts = []
                for cell, cached_cell in zip(row, cached_row, strict=True):
                    value = cell.value
                    if cell.data_type == 'f':
                        formulas += 1
                        if cached_cell.value is None:
                            missing_results += 1
                            rendered = '[formula; cached value unavailable; do not infer a value]'
                        else:
                            rendered = f'{_cell_text(cached_cell.value)} [cached formula value; may be stale]'
                    elif value is None or (isinstance(value, str) and not value.strip()):
                        continue
                    else:
                        rendered = _cell_text(value)
                    values += 1
                    row_parts.append(f'{cell.coordinate}={rendered}')
                if row_parts:
                    append(f'Row {row_index}: ' + ' | '.join(row_parts))
        if not values:
            raise AIImportError('The Excel workbook is empty or contains no readable cell values.')
        warnings = []
        if formulas:
            warnings.append(f'{formulas} formula cell(s) were not evaluated; saved results may be stale.')
        if missing_results:
            warnings.append(
                f'{missing_results} formula cell(s) have no cached value. Recalculate and save in Excel/LibreOffice, '
                'or replace formulas with values before importing.'
            )
        return '\n'.join(parts), warnings
    except AIImportError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise AIImportError('Unable to read Excel workbook: malformed, encrypted, or unsupported .xlsx file.') from exc
    finally:
        for workbook in workbooks:
            workbook.close()
