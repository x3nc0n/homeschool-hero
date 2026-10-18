from __future__ import annotations

import csv
import calendar
import hashlib
import math
import re
from datetime import UTC, date, datetime, time, timedelta
from io import BytesIO, StringIO
from typing import Any

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from backend.models import Assignment, LessonPlan, Student, Subject
from backend.models.calendar import SchoolYear
from backend.models.curriculum import (
    CurriculumLesson,
    CurriculumPackage,
    CurriculumUnit,
    LessonResource,
    Resource,
    ResourceType,
)
from backend.schemas.bulk_assignment_import import (
    BulkAssignmentImportItem,
    BulkAssignmentImportTarget,
    ParsedAssignmentCandidate,
    ParsedAssignmentDocument,
    WeeklyImportContext,
    WeeklyImportItem,
)
from backend.services.assignment_spreadsheet import XLSX_CONTENT_TYPE
from backend.services.curriculum_ai_import import AIImportError

WEEKLY_HEADERS = (
    'student', 'grade', 'subject', 'unit', 'lesson', 'instructions', 'date', 'day',
    'curriculum', 'category', 'estimated_minutes', 'resource_url', 'max_score', 'weight',
)
REQUIRED_HEADERS = set(WEEKLY_HEADERS[:6])
TEMPLATE_GUIDANCE = (
    'One row per student and daily activity. Required: student, grade, subject, unit, lesson, instructions, '
    'and at least one of date or day. Dates must be YYYY-MM-DD. Weekdays resolve against the selected Monday; '
    'when both are present they must agree. Replace the example student with an existing student. '
    'Optional: curriculum, category (homework, quiz, test, project, participation, extra_credit, other), '
    'estimated_minutes (1-1440), resource_url (http/https), max_score (>0), weight (>=0). '
    'Formulas are not evaluated; replace them with literal values. No records are created until confirmation.'
    ' XLSX subject-by-weekday grids are also supported when they have a Subject header and weekday columns. '
    'Blank assignments and checkmark columns are ignored. A dated title must agree with the selected week; '
    'a title without a year uses the selected week context, never the current date. Missing grades require clarification.'
)


def normalize(value: Any) -> str:
    return ' '.join(str(value or '').casefold().split())


def template_bytes(format: str) -> bytes:
    row = ['Lucas', '7', 'Spelling', 'Unit 1', 'Lesson 6', 'Page 25', '', 'Tuesday',
           '', 'homework', '20', '', '100', '1']
    if format == 'csv':
        stream = StringIO(newline='')
        writer = csv.writer(stream)
        writer.writerow(WEEKLY_HEADERS)
        writer.writerow(row)
        return stream.getvalue().encode('utf-8')
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Weekly plan'
    sheet.append(WEEKLY_HEADERS)
    sheet.append(row)
    sheet.freeze_panes = 'A2'
    guidance = workbook.create_sheet('Instructions')
    guidance.append(['Weekly planning template'])
    guidance.append([TEMPLATE_GUIDANCE])
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def _strict_date(value: Any) -> date | None:
    if value is None or value == '':
        return None
    if isinstance(value, datetime):
        if value.time() != time.min:
            raise ValueError('date cells must not contain a time')
        return value.date()
    if isinstance(value, date):
        return value
    if not re.fullmatch(r'\d{4}-\d{2}-\d{2}', str(value)):
        raise ValueError('date must be YYYY-MM-DD; ambiguous dates are not guessed')
    return date.fromisoformat(str(value))


def scheduled_date(weekly: WeeklyImportItem, context: WeeklyImportContext) -> date | None:
    day_date = None
    if weekly.day:
        weekdays = ('monday', 'tuesday', 'wednesday', 'thursday', 'friday', 'saturday', 'sunday')
        day = normalize(weekly.day)
        matches = [index for index, name in enumerate(weekdays) if day in {name, name[:3]}]
        if not matches:
            raise ValueError('day must be a weekday name, such as Tuesday')
        day_date = context.week_start + timedelta(days=matches[0])
    explicit = weekly.scheduled_date
    if explicit and day_date and explicit != day_date:
        raise ValueError('date and day conflict; correct one before confirming')
    result = explicit or day_date
    if result and not context.week_start <= result <= context.week_start + timedelta(days=6):
        raise ValueError('date is outside the selected week; change the date or selected week')
    return result


def parse_template(source, *, max_items: int) -> ParsedAssignmentDocument:
    """Read only after the existing bounded upload extraction/preflight has succeeded."""
    tables = []
    if source.content_type in {'text/csv', 'text/tab-separated-values'}:
        delimiter = '\t' if source.content_type == 'text/tab-separated-values' else ','
        tables.append((source.filename, list(csv.reader(StringIO(source.extracted.text), delimiter=delimiter))))
    elif source.content_type == XLSX_CONTENT_TYPE:
        workbook = load_workbook(BytesIO(source.raw_bytes), read_only=True, data_only=False, keep_links=False)
        try:
            for sheet in workbook.worksheets:
                if sheet.title == 'Instructions':
                    continue
                sheet.reset_dimensions()
                table = []
                for row in sheet.iter_rows():
                    if any(cell.data_type == 'f' for cell in row):
                        raise AIImportError(f'{sheet.title}: formulas must be replaced with literal values')
                    table.append([cell.value for cell in row])
                tables.append((sheet.title, table))
        finally:
            workbook.close()
    else:
        raise AIImportError('Weekly templates must be CSV, TSV, or XLSX')
    candidates = []
    for sheet, table in tables:
        if not table or not any(table[0]):
            raise AIImportError(f'{sheet}: missing canonical header row')
        headers = [normalize(value) for value in table[0]]
        while headers and not headers[-1]:
            headers.pop()
        if len(set(headers)) != len(headers):
            raise AIImportError(f'{sheet}: duplicate headers are not allowed')
        missing = REQUIRED_HEADERS - set(headers)
        if missing or not {'date', 'day'} & set(headers):
            raise AIImportError(f'{sheet}: required headers: {", ".join(WEEKLY_HEADERS[:8])}; missing {sorted(missing)}')
        unknown = set(headers) - set(WEEKLY_HEADERS)
        if unknown:
            raise AIImportError(f'{sheet}: unknown headers: {", ".join(sorted(unknown))}')
        for row_number, values in enumerate(table[1:], 2):
            if not any(value is not None and str(value).strip() for value in values):
                continue
            if len(values) > len(headers) and any(value not in (None, '') for value in values[len(headers):]):
                raise AIImportError(f'{sheet} row {row_number}: more values than headers')
            row = dict(zip(headers, values))
            row = {key: value.strip() if isinstance(value, str) else value for key, value in row.items()}
            try:
                weekly = WeeklyImportItem(
                    **{key: str(row.get(key) or '') for key in WEEKLY_HEADERS[:6]},
                    curriculum=row.get('curriculum') or None,
                    scheduled_date=_strict_date(row.get('date')),
                    day=str(row.get('day') or '') or None,
                    estimated_minutes=row.get('estimated_minutes') if row.get('estimated_minutes') not in (None, '') else None,
                    resource_url=row.get('resource_url') or None,
                    source_location=f'{sheet} row {row_number}',
                )
                candidate = ParsedAssignmentCandidate(
                    weekly=weekly,
                    title=weekly.lesson or None,
                    description=weekly.instructions or None,
                    subject_ref=weekly.subject or None,
                    student_refs=[weekly.student] if weekly.student else [],
                    category=row.get('category') or None,
                    weight=row.get('weight') if row.get('weight') not in (None, '') else None,
                    max_score=row.get('max_score') if row.get('max_score') not in (None, '') else None,
                )
                if any(value is not None and not math.isfinite(value) for value in (candidate.weight, candidate.max_score)):
                    raise ValueError('weight and max_score must be finite numbers')
            except (ValueError, ValidationError) as exc:
                raise AIImportError(f'{sheet} row {row_number}: {exc}') from exc
            candidates.append(candidate)
            if len(candidates) > max_items:
                raise AIImportError(f'Weekly template exceeds {max_items} items; split the file')
    if not candidates:
        raise AIImportError('Weekly template contains no plan rows')
    return ParsedAssignmentDocument(assignments=candidates)


def parse_weekday_grid(source, *, context: WeeklyImportContext, max_items: int) -> ParsedAssignmentDocument | None:
    if source.content_type != XLSX_CONTENT_TYPE:
        return None
    workbook = load_workbook(BytesIO(source.raw_bytes), read_only=True, data_only=False, keep_links=False)
    candidates = []
    recognized = False
    unknown_sheets = []
    weekdays = {name.casefold(): index for index, name in enumerate(calendar.day_name)}
    try:
        for sheet in workbook.worksheets:
            if sheet.title == 'Instructions':
                continue
            sheet.reset_dimensions()
            rows = list(sheet.iter_rows())
            header = None
            for index, cells in enumerate(rows[:20]):
                values = [normalize(cell.value) for cell in cells]
                if 'subject' in values and sum(value in weekdays for value in values) >= 2:
                    header = (index, values.index('subject'), {
                        column: weekdays[value] for column, value in enumerate(values) if value in weekdays
                    })
                    if len(header[2].values()) != len(set(header[2].values())):
                        raise AIImportError(f'{sheet.title}: duplicate weekday columns')
                    break
            if header is None:
                if any(cell.value is not None for cells in rows for cell in cells):
                    unknown_sheets.append(sheet.title)
                continue
            recognized = True
            header_index, subject_column, day_columns = header
            title_parts = [str(cell.value).strip() for cells in rows[:header_index] for cell in cells
                           if cell.value is not None and str(cell.value).strip()]
            title = ' '.join(title_parts)
            student_match = re.search(r'^(.*?)\s+Week\s+\d+\b', title, re.I)
            student_ref = student_match[1].strip() if student_match else ''
            for row_index, cells in enumerate(rows[header_index + 1:], header_index + 2):
                if any(cell.data_type == 'f' for cell in cells):
                    raise AIImportError(f'{sheet.title} row {row_index}: formulas must be replaced with literal values')
                subject = str(cells[subject_column].value or '').strip() if subject_column < len(cells) else ''
                for column, weekday in day_columns.items():
                    if column >= len(cells):
                        continue
                    cell = cells[column]
                    value = cell.value
                    if value is None or str(value).strip() in {'', '✓', '✔', '☑'}:
                        continue
                    if not subject:
                        raise AIImportError(f'{sheet.title} {cell.coordinate}: assignment has no subject label')
                    if not isinstance(value, str):
                        raise AIImportError(f'{sheet.title} {cell.coordinate}: assignment must contain text')
                    lines = [line.strip() for line in value.splitlines() if line.strip()]
                    heading = lines[0]
                    unit_match = re.match(r'^(Unit\s+\S+)\s*[—–-]\s*(.+)$', heading, re.I)
                    unit = unit_match[1].strip() if unit_match else heading
                    lesson = unit_match[2].strip() if unit_match else heading
                    instructions = '\n'.join(lines[1:]) or lesson
                    try:
                        weekly = WeeklyImportItem(
                            student=student_ref, grade='', subject=subject, unit=unit, lesson=lesson,
                            instructions=instructions, day=calendar.day_name[weekday],
                            source_location=f'{sheet.title} {cell.coordinate}', source_week_label=title[:255] or None,
                        )
                        candidates.append(ParsedAssignmentCandidate(
                            weekly=weekly, title=lesson, description=instructions,
                            student_refs=[student_ref] if student_ref else [], subject_ref=subject,
                            source_excerpt=value[:1000],
                            category='quiz' if re.search(r'\bquiz\b', lesson, re.I) else None,
                            notes=[] if unit_match else [
                                'No separate unit heading was supplied; the source lesson/chapter heading is used as the curriculum grouping.',
                            ],
                        ))
                    except ValidationError as exc:
                        raise AIImportError(f'{sheet.title} {cell.coordinate}: {exc}') from exc
                    if len(candidates) > max_items:
                        raise AIImportError(f'Weekly grid exceeds {max_items} items; split the file')
        if recognized and unknown_sheets:
            raise AIImportError(f'{", ".join(unknown_sheets)}: expected another subject-by-weekday grid')
        if recognized and not candidates:
            raise AIImportError('Weekly grid contains no daily assignments')
        return ParsedAssignmentDocument(assignments=candidates) if recognized else None
    finally:
        workbook.close()


def validate_source_week(label: str | None, context: WeeklyImportContext) -> None:
    if not label:
        return
    months = {name.casefold(): index for index, name in enumerate(calendar.month_name) if name}
    pattern = r'\b(' + '|'.join(months) + r')\s+(\d{1,2})\s*[-–—]\s*(\d{1,2})(?:,?\s+(\d{4}))?\b'
    match = re.search(pattern, label, re.I)
    if match is None:
        if re.search(r'\b(?:' + '|'.join(months) + r')\b|\d{4}-\d{2}-\d{2}', label, re.I):
            raise ValueError('Grid title contains an unrecognized explicit date; use the canonical template to preserve it')
        return
    month = months[match[1].casefold()]
    year = int(match[4]) if match[4] else context.week_start.year
    start = date(year, month, int(match[2]))
    end = date(year, month, int(match[3]))
    if start > end or not (context.week_start <= start <= end <= context.week_start + timedelta(days=6)):
        raise ValueError(f'Grid title "{label}" conflicts with selected week {context.week_start}; select its week and upload again')


def qualified_subject(subject: str, grade: str) -> str:
    grade = grade.strip()
    match = re.fullmatch(r'(?:grade\s*)?(\d{1,2})(?:st|nd|rd|th)?(?:\s*grade)?', grade, re.I)
    if match:
        number = int(match[1])
        suffix = 'th' if 10 <= number % 100 <= 20 else {1: 'st', 2: 'nd', 3: 'rd'}.get(number % 10, 'th')
        grade = f'{number}{suffix} Grade'
    elif normalize(grade) in {'k', 'kindergarten'}:
        grade = 'Kindergarten'
    subject = subject.strip()
    base = re.sub(r'\s*-\s*(?:(?:grade\s*)?\d{1,2}(?:st|nd|rd|th)?(?:\s*grade)?|kindergarten)$',
                  '', subject, flags=re.I).strip()
    name = f'{base} - {grade}'
    if base != subject and normalize(name) != normalize(subject):
        raise ValueError('Subject grade suffix conflicts with the grade column')
    if len(name) > 120:
        raise ValueError('Grade-qualified subject exceeds 120 characters')
    return name


def activity_name(weekly: WeeklyImportItem) -> str:
    name = f'{weekly.lesson.strip()} - {" ".join(weekly.instructions.split())}'
    if len(name) <= 160:
        return name
    digest = hashlib.sha256(normalize(name).encode()).hexdigest()[:12]
    return f'{name[:143]} - {digest}'


async def validate_context(db: AsyncSession, *, family_id: int, context: WeeklyImportContext) -> SchoolYear:
    year = (await db.execute(select(SchoolYear).where(
        SchoolYear.id == context.school_year_id, SchoolYear.family_id == family_id,
    ))).scalar_one_or_none()
    if year is None:
        raise HTTPException(404, 'School year not found')
    return year


async def resolve_items(
    db: AsyncSession, *, family_id: int, items: list[BulkAssignmentImportItem], context: WeeklyImportContext,
) -> list[dict]:
    await validate_context(db, family_id=family_id, context=context)
    students = (await db.execute(select(Student).where(Student.family_id == family_id))).scalars().all()
    subjects = (await db.execute(select(Subject).where(Subject.family_id == family_id))).scalars().all()
    packages = (await db.execute(select(CurriculumPackage).where(
        CurriculumPackage.family_id == family_id, CurriculumPackage.school_year_id == context.school_year_id,
    ))).scalars().all()
    units = (await db.execute(select(CurriculumUnit).join(CurriculumPackage).where(
        CurriculumPackage.family_id == family_id, CurriculumPackage.school_year_id == context.school_year_id,
    ))).scalars().all()
    lessons = (await db.execute(select(CurriculumLesson).join(CurriculumUnit).join(CurriculumPackage).where(
        CurriculumPackage.family_id == family_id, CurriculumPackage.school_year_id == context.school_year_id,
    ))).scalars().all()
    plans = (await db.execute(select(LessonPlan).where(LessonPlan.family_id == family_id))).scalars().all()
    assignments = (await db.execute(select(Assignment).where(
        Assignment.family_id == family_id, Assignment.lesson_plan_id.is_not(None),
    ))).scalars().all()
    resource_links = (await db.execute(
        select(LessonResource.lesson_id, Resource.url).join(Resource, Resource.id == LessonResource.resource_id)
        .where(Resource.family_id == family_id)
    )).all()
    questions = []
    seen: dict[tuple, tuple[str, date | None, tuple]] = {}
    for index, item in enumerate(items):
        weekly = item.weekly
        item.errors = []
        item.missing_fields = []
        item.ambiguous_fields = []
        if weekly is None:
            item.errors.append('Weekly metadata is required')
            item.status = 'invalid'
            continue
        resolution: dict[str, Any] = {}
        weekly.resolution = resolution
        for field in ('student', 'grade', 'subject', 'unit', 'lesson', 'instructions'):
            if not getattr(weekly, field).strip():
                item.missing_fields.append(f'weekly.{field}')
        try:
            validate_source_week(weekly.source_week_label, context)
            resolved_date = scheduled_date(weekly, context)
            if resolved_date is None:
                item.missing_fields.append('weekly.scheduled_date')
            elif item.due_date is not None:
                # An explicit assignment-date edit also schedules the linked lesson plan.
                edited_date = item.due_date.date()
                original = item.refs.get('weekly_original_due_date')
                if original and item.due_date.isoformat() != original:
                    weekly.scheduled_date = edited_date
                    resolved_date = scheduled_date(weekly, context)
            if resolved_date and item.due_date is None:
                item.due_date = datetime.combine(resolved_date, time.min, tzinfo=UTC)
            elif resolved_date and item.due_date and (
                not item.refs.get('weekly_original_due_date')
                or item.due_date.isoformat() == item.refs['weekly_original_due_date']
            ):
                item.due_date = datetime.combine(resolved_date, time.min, tzinfo=UTC)
            if resolved_date:
                weekly.scheduled_date = resolved_date
                item.refs['weekly_original_due_date'] = item.due_date.isoformat() if item.due_date else None
            expected_subject = qualified_subject(weekly.subject, weekly.grade)
            weekly.activity_name = activity_name(weekly)
        except ValueError as exc:
            item.errors.append(str(exc))
            expected_subject = ''
        if not item.title:
            item.missing_fields.append('title')
        if item.recurrence.value != 'none':
            item.errors.append('Weekly activities cannot recur; import each dated activity separately')
        if not math.isfinite(item.weight) or not math.isfinite(item.max_score):
            item.errors.append('weight and max_score must be finite numbers')
        prior_student_ref = item.refs.get('weekly_original_student_ref')
        if (prior_student_ref is not None and normalize(prior_student_ref) != normalize(weekly.student)
                and [target.student_id for target in item.targets] == item.refs.get('weekly_original_target_ids')):
            item.targets = []
        if not item.targets:
            matches = [record for record in students if normalize(record.name) == normalize(weekly.student)]
            if len(matches) == 1:
                item.targets = [BulkAssignmentImportTarget(student_id=matches[0].id)]
            else:
                item.missing_fields.append('targets')
                resolution['student'] = 'ambiguous' if matches else 'missing'
                questions.append({
                    'id': f'weekly_student_{index}', 'field': 'targets', 'assignment_indexes': [index],
                    'message': f'Choose the existing student for {weekly.student or "this row"}. Students are never created.',
                    'choices': [{'id': record.id, 'label': record.name} for record in matches or students],
                    'allow_apply_to_all': True,
                })
        if len(item.targets) > 1:
            item.errors.append('Weekly rows must target exactly one student')
        if any(target.due_date and target.due_date.date() != weekly.scheduled_date for target in item.targets):
            item.errors.append('Student due date must agree with the scheduled activity date')
        if item.targets:
            student = next((record for record in students if record.id == item.targets[0].student_id), None)
            if student is None:
                item.errors.append('Student not found in this family')
            else:
                resolution.update(student='reused', student_id=student.id)
        item.refs['weekly_original_student_ref'] = weekly.student
        item.refs['weekly_original_target_ids'] = [target.student_id for target in item.targets]
        matches = [record for record in subjects if normalize(record.name) == normalize(expected_subject)]
        if (item.refs.get('weekly_original_subject_name') is not None
                and normalize(item.refs['weekly_original_subject_name']) != normalize(expected_subject)
                and item.subject_id == item.refs.get('weekly_original_subject_id')):
            item.subject_id = None
        if item.subject_id:
            selected = next((record for record in subjects if record.id == item.subject_id), None)
            if selected is None or normalize(selected.name) != normalize(expected_subject):
                item.errors.append(f'Subject must be grade-qualified as {expected_subject}')
        elif len(matches) == 1:
            item.subject_id = matches[0].id
        elif len(matches) > 1:
            item.ambiguous_fields.append('subject_id')
            questions.append({
                'id': f'weekly_subject_{index}', 'field': 'subject_id', 'assignment_indexes': [index],
                'message': f'Choose the existing grade-qualified subject {expected_subject}.',
                'choices': [{'id': record.id, 'label': record.name} for record in matches],
            })
        resolution.update(subject='reused' if item.subject_id else ('ambiguous' if len(matches) > 1 else 'new'),
                          subject_id=item.subject_id, subject_name=expected_subject)
        item.refs['weekly_original_subject_name'] = expected_subject
        item.refs['weekly_original_subject_id'] = item.subject_id
        candidates = [record for record in packages if record.subject_id == item.subject_id] if item.subject_id else []
        if weekly.curriculum:
            candidates = [record for record in candidates if normalize(record.name) == normalize(weekly.curriculum)]
        package = None
        if weekly.package_id:
            package = next((record for record in candidates if record.id == weekly.package_id), None)
            if package is None:
                item.errors.append('Selected curriculum does not match this subject, family, and school year')
        elif len(candidates) == 1:
            package = candidates[0]
        elif len(candidates) > 1:
            item.ambiguous_fields.append('weekly.package_id')
            questions.append({
                'id': f'weekly_package_{index}', 'field': 'weekly.package_id', 'assignment_indexes': [index],
                'message': 'Choose a curriculum package for this subject.',
                'choices': [{'id': record.id, 'label': record.name} for record in candidates],
                'allow_apply_to_all': True,
            })
        proposed_package_name = weekly.curriculum or f'{expected_subject} Curriculum'
        if len(proposed_package_name) > 160:
            item.errors.append('Curriculum name exceeds 160 characters')
        collision = next((record for record in packages if normalize(record.name) == normalize(proposed_package_name)
                          and record.subject_id != item.subject_id), None)
        if not package and collision:
            item.errors.append('Curriculum name belongs to another subject; supply a distinct curriculum name')
        resolution.update(
            package='reused' if package else ('ambiguous' if len(candidates) > 1 else 'new'),
            package_id=package.id if package else None,
            package_name=package.name if package else proposed_package_name,
            package_choices=[{'id': record.id, 'label': record.name} for record in candidates],
        )
        matching_units = [record for record in units if package and record.package_id == package.id
                          and normalize(record.name) == normalize(weekly.unit)]
        unit = matching_units[0] if len(matching_units) == 1 else None
        if len(matching_units) > 1:
            item.errors.append('Multiple matching curriculum units; rename the unit to disambiguate')
        matching_lessons = [record for record in lessons if unit and record.unit_id == unit.id
                            and normalize(record.name) in {normalize(weekly.activity_name), normalize(weekly.lesson)}
                            and normalize(record.description) == normalize(weekly.instructions)]
        lesson = matching_lessons[0] if len(matching_lessons) == 1 else None
        if len(matching_lessons) > 1:
            item.errors.append('Multiple matching curriculum activities; edit lesson or instructions to disambiguate')
        if unit and not lesson and any(record.unit_id == unit.id and normalize(record.name) == normalize(weekly.activity_name)
                                      for record in lessons):
            item.errors.append('An activity with this name has different instructions; edit lesson or instructions')
        resolution.update(unit='reused' if unit else 'new', unit_id=unit.id if unit else None,
                          activity='reused' if lesson else 'new', activity_id=lesson.id if lesson else None)
        student_id = item.targets[0].student_id if len(item.targets) == 1 else None
        plan = next((record for record in plans if lesson and record.student_id == student_id
                     and record.curriculum_lesson_id == lesson.id), None)
        resolution.update(lesson_plan='reused' if plan else 'new', lesson_plan_id=plan.id if plan else None)
        if plan and (plan.target_date != weekly.scheduled_date or plan.school_year_id != context.school_year_id):
            resolution['lesson_plan'] = 'conflict'
            item.errors.append(f'Existing lesson plan {plan.id} is scheduled for {plan.target_date}; it will not be moved')
        if plan and weekly.estimated_minutes is not None and plan.estimated_duration_minutes != weekly.estimated_minutes:
            item.errors.append('Existing lesson plan has different estimated minutes; it will not be overwritten')
        existing_assignment = next((record for record in assignments if plan and record.lesson_plan_id == plan.id), None)
        resolution.update(assignment='reused' if existing_assignment else 'new',
                          assignment_id=existing_assignment.id if existing_assignment else None)
        if existing_assignment:
            effective_date = existing_assignment.due_date.date() if existing_assignment.due_date else None
            if (normalize(existing_assignment.title) != normalize(item.title)
                    or normalize(existing_assignment.description) != normalize(item.description)
                    or effective_date != (item.due_date.date() if item.due_date else None)
                    or existing_assignment.category != item.category
                    or existing_assignment.max_score != item.max_score or existing_assignment.weight != item.weight
                    or existing_assignment.subject_id != item.subject_id
                    or {target.student_id for target in existing_assignment.targets} != {student_id}
                    or any(target.due_date and target.due_date.date() != weekly.scheduled_date
                           for target in existing_assignment.targets)
                    or existing_assignment.grading_period_id != item.grading_period_id):
                resolution['assignment'] = 'conflict'
                item.errors.append('Existing assignment differs from this row; it will not be overwritten')
            elif weekly.resource_url and (lesson.id, weekly.resource_url) not in resource_links:
                resolution['assignment'] = 'conflict'
                item.errors.append('Existing activity has different resource links; it will not be overwritten')
            else:
                resolution['duplicate'] = True
        identity = (student_id, normalize(expected_subject), package.id if package else normalize(proposed_package_name),
                    normalize(weekly.unit), lesson.id if lesson else normalize(weekly.activity_name))
        signature = (normalize(item.title), normalize(item.description), item.category, item.max_score, item.weight,
                     item.grading_period_id, weekly.estimated_minutes, weekly.resource_url)
        if identity in seen:
            prior_id, prior_date, prior_signature = seen[identity]
            if prior_date != weekly.scheduled_date:
                item.errors.append('The same student/activity appears on conflicting dates')
            elif prior_signature != signature:
                item.errors.append('Duplicate activity rows have different assignment fields; correct or remove one row')
            else:
                resolution['duplicate_of'] = prior_id
                resolution['duplicate'] = True
        else:
            seen[identity] = (item.client_item_id, weekly.scheduled_date, signature)
        item.status = ('invalid' if item.errors else
                       'needs_clarification' if item.missing_fields or item.ambiguous_fields else 'ready')
    grade_indexes = [index for index, item in enumerate(items) if 'weekly.grade' in item.missing_fields]
    if grade_indexes:
        questions.append({
            'id': 'weekly_grade', 'field': 'weekly.grade', 'assignment_indexes': grade_indexes,
            'message': 'Which grade should be used? Student records do not store grade levels.',
            'choices': [], 'allow_apply_to_all': True,
        })
    return questions


async def create_links(
    db: AsyncSession, *, family_id: int, user_id: int, items: list[BulkAssignmentImportItem],
    context: WeeklyImportContext,
) -> tuple[list[BulkAssignmentImportItem], list[str], dict]:
    counts = {key: {'created': 0, 'reused': 0} for key in ('subjects', 'packages', 'units', 'activities')}
    curriculum_ids = {key: {'created': [], 'reused': []} for key in counts}
    counted = {key: set() for key in counts}
    created_plans = []
    reused_plans = set()
    reused_assignments = set()
    ready = []
    skipped = []
    seen = set()
    for item in items:
        # Re-resolve after every row so new sibling records are reused within this transaction.
        await resolve_items(db, family_id=family_id, items=[item], context=context)
        if item.status != 'ready':
            raise HTTPException(409, {'code': 'invalid_import_items', 'items': [item.model_dump(mode='json')]})
        weekly = item.weekly
        resolution = weekly.resolution
        subject = await db.get(Subject, item.subject_id) if item.subject_id else None
        if subject is None:
            subject = Subject(family_id=family_id, name=resolution['subject_name'])
            db.add(subject)
            await db.flush()
        item.subject_id = subject.id
        package_id = resolution['package_id']
        package = await db.get(CurriculumPackage, package_id) if package_id else None
        if package is None:
            package = CurriculumPackage(
                family_id=family_id, school_year_id=context.school_year_id, subject_id=subject.id,
                name=resolution['package_name'], created_by_user_id=user_id,
            )
            db.add(package)
            await db.flush()
        unit_id = resolution['unit_id']
        unit = await db.get(CurriculumUnit, unit_id) if unit_id else None
        if unit is None:
            unit = CurriculumUnit(package_id=package.id, name=weekly.unit, standards_tags=[])
            db.add(unit)
            await db.flush()
        lesson_id = resolution['activity_id']
        lesson = await db.get(CurriculumLesson, lesson_id) if lesson_id else None
        if lesson is None:
            lesson = CurriculumLesson(
                unit_id=unit.id, name=weekly.activity_name, description=weekly.instructions,
                estimated_duration_minutes=weekly.estimated_minutes, standards_tags=[],
            )
            db.add(lesson)
            await db.flush()
        for key, record, decision in (
            ('subjects', subject, resolution['subject']), ('packages', package, resolution['package']),
            ('units', unit, resolution['unit']), ('activities', lesson, resolution['activity']),
        ):
            if record.id not in counted[key]:
                operation = 'reused' if decision == 'reused' else 'created'
                counts[key][operation] += 1
                curriculum_ids[key][operation].append(record.id)
                counted[key].add(record.id)
        student_id = item.targets[0].student_id
        identity = (student_id, lesson.id)
        if resolution.get('lesson_plan_id') and resolution['lesson_plan_id'] not in created_plans:
            reused_plans.add(resolution['lesson_plan_id'])
        if resolution.get('assignment_id'):
            reused_assignments.add(resolution['assignment_id'])
        if identity in seen or resolution.get('duplicate'):
            skipped.append(item.client_item_id)
            continue
        seen.add(identity)
        plan_id = resolution['lesson_plan_id']
        plan = await db.get(LessonPlan, plan_id) if plan_id else None
        if plan is None:
            plan = LessonPlan(
                family_id=family_id, student_id=student_id, curriculum_lesson_id=lesson.id,
                school_year_id=context.school_year_id, target_date=weekly.scheduled_date,
                estimated_duration_minutes=weekly.estimated_minutes, notes=weekly.instructions,
            )
            db.add(plan)
            await db.flush()
            created_plans.append(plan.id)
        if weekly.resource_url:
            digest = hashlib.sha256(weekly.resource_url.encode()).hexdigest()[:16]
            name = f'Lesson resource {digest}'
            resource = (await db.execute(select(Resource).where(
                Resource.family_id == family_id, Resource.name == name,
            ))).scalar_one_or_none()
            if resource is None:
                resource = Resource(
                    family_id=family_id, name=name, resource_type=ResourceType.link, url=weekly.resource_url,
                    created_by_user_id=user_id, tags=[], resource_metadata={},
                )
                db.add(resource)
                await db.flush()
            link = await db.get(LessonResource, (lesson.id, resource.id))
            if link is None:
                db.add(LessonResource(lesson_id=lesson.id, resource_id=resource.id))
                await db.flush()
        item.lesson_plan_id = plan.id
        ready.append(item)
    return ready, skipped, {
        'created_lesson_plan_ids': created_plans,
        'reused_lesson_plan_ids': sorted(reused_plans),
        'reused_assignment_ids': sorted(reused_assignments),
        'lesson_plan_count': len(created_plans),
        'curriculum_counts': counts,
        'curriculum_ids': curriculum_ids,
    }
